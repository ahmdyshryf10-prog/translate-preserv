import os
import io
import json
import csv
import html
import time
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

import streamlit as st

# Optional PDF
try:
    import fitz
except ImportError:
    fitz = None

# Optional DOCX
try:
    from docx import Document
except ImportError:
    Document = None

# Gemini
try:
    from google import genai
    from google.genai import types
except ImportError:
    genai = None
    types = None

# OpenAI
try:
    from openai import OpenAI
except ImportError:
    OpenAI = None


# =========================================================
# CONFIG
# =========================================================

st.set_page_config(
    page_title="Translate Preserv",
    page_icon="✦",
    layout="wide",
    initial_sidebar_state="collapsed",
)


# =========================================================
# THEME
# =========================================================

st.markdown(
    """
    <style>

    :root {
        --bg: #071111;
        --panel: #0d1b1b;
        --panel2: #102222;
        --border: #1c3838;
        --text: #edfafa;
        --muted: #8daaaa;
        --accent: #36d6c0;
        --accent2: #20a99a;
    }

    .stApp {
        background:
            radial-gradient(
                circle at 50% -10%,
                rgba(54,214,192,.13),
                transparent 35%
            ),
            var(--bg);
        color: var(--text);
    }

    .block-container {
        max-width: 1050px;
        padding-top: 2rem;
        padding-bottom: 5rem;
    }

    .hero {
        text-align: center;
        padding: 25px 0 20px 0;
    }

    .hero-icon {
        font-size: 34px;
        margin-bottom: 8px;
    }

    .hero-title {
        font-size: 38px;
        font-weight: 750;
        letter-spacing: -1px;
    }

    .hero-subtitle {
        color: var(--muted);
        margin-top: 7px;
        font-size: 15px;
    }

    .glass {
        background: rgba(13,27,27,.78);
        border: 1px solid var(--border);
        border-radius: 22px;
        padding: 20px;
        box-shadow: 0 12px 45px rgba(0,0,0,.18);
    }

    .status {
        display: inline-flex;
        align-items: center;
        gap: 7px;
        background: rgba(54,214,192,.09);
        border: 1px solid rgba(54,214,192,.20);
        color: var(--accent);
        padding: 6px 11px;
        border-radius: 999px;
        font-size: 12px;
    }

    .dot {
        width: 7px;
        height: 7px;
        background: var(--accent);
        border-radius: 50%;
        display: inline-block;
    }

    div[data-testid="stFileUploader"] {
        border-radius: 18px;
    }

    .small-muted {
        color: var(--muted);
        font-size: 12px;
    }

    button[kind="primary"] {
        border-radius: 14px !important;
    }

    textarea {
        border-radius: 16px !important;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# HEADER
# =========================================================

st.markdown(
    """
    <div class="hero">
        <div class="hero-icon">✦</div>
        <div class="hero-title">Translate Preserv</div>
        <div class="hero-subtitle">
            Translate text and documents while preserving structure.
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# SECRETS
# =========================================================

def secret(name, default=None):
    try:
        value = st.secrets.get(name)
        if value is not None:
            return value
    except Exception:
        pass

    return os.getenv(name, default)


GEMINI_API_KEY = secret("GEMINI_API_KEY")
GEMINI_MODEL = secret(
    "GEMINI_MODEL",
    "gemini-3.8-flash",
)

OPENAI_API_KEY = secret("OPENAI_API_KEY")
OPENAI_MODEL = secret(
    "OPENAI_MODEL",
    "gpt-5.6-luna",
)


# =========================================================
# PROVIDERS
# =========================================================

PROVIDERS = {}

if GEMINI_API_KEY:
    PROVIDERS["Gemini"] = {
        "type": "gemini",
        "model": GEMINI_MODEL,
        "key": GEMINI_API_KEY,
    }

if OPENAI_API_KEY:
    PROVIDERS["OpenAI"] = {
        "type": "openai",
        "model": OPENAI_MODEL,
        "key": OPENAI_API_KEY,
    }


# =========================================================
# LANGUAGES
# =========================================================

LANGUAGES = [
    "Persian",
    "English",
    "Arabic",
    "German",
    "French",
    "Spanish",
    "Italian",
    "Portuguese",
    "Dutch",
    "Russian",
    "Ukrainian",
    "Polish",
    "Czech",
    "Slovak",
    "Romanian",
    "Hungarian",
    "Greek",
    "Turkish",
    "Hebrew",
    "Urdu",
    "Hindi",
    "Bengali",
    "Indonesian",
    "Malay",
    "Vietnamese",
    "Thai",
    "Chinese Simplified",
    "Chinese Traditional",
    "Japanese",
    "Korean",
    "Swedish",
    "Norwegian",
    "Danish",
    "Finnish",
    "Pashto",
    "Kurdish",
    "Azerbaijani",
    "Armenian",
    "Georgian",
]


RTL_LANGUAGES = {
    "Persian",
    "Arabic",
    "Hebrew",
    "Urdu",
    "Pashto",
    "Kurdish",
}


def is_rtl(language):
    return language in RTL_LANGUAGES


# =========================================================
# SYSTEM PROMPT
# =========================================================

SYSTEM_PROMPT = """
You are a professional document translation engine.

Translate faithfully and completely.

Rules:

1. Never summarize.
2. Never explain the translation.
3. Never add commentary.
4. Never intentionally remove information.
5. Preserve names, numbers, dates and technical terms.
6. Preserve paragraph structure.
7. Preserve the order of content.
8. Preserve placeholders.
9. Preserve URLs, emails and file names.
10. Do not invent missing content.
11. Preserve the author's tone and style.
12. Keep terminology consistent throughout the document.
13. Do not translate proper names unless appropriate for the target language.
14. Return only the translation.
"""


# =========================================================
# CLIENTS
# =========================================================

@st.cache_resource
def get_gemini_client(api_key):
    if genai is None:
        raise RuntimeError(
            "google-genai is not installed."
        )

    return genai.Client(api_key=api_key)


@st.cache_resource
def get_openai_client(api_key):
    if OpenAI is None:
        raise RuntimeError(
            "openai is not installed."
        )

    return OpenAI(api_key=api_key)


# =========================================================
# GEMINI
# =========================================================

def translate_gemini(
    text,
    target_language,
    model,
    api_key,
):

    client = get_gemini_client(api_key)

    prompt = f"""
Target language: {target_language}

Translate the following text.

TEXT:
{text}
"""

    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.2,
        ),
    )

    result = response.text

    if not result:
        raise RuntimeError(
            "Gemini returned an empty response."
        )

    return result.strip()


# =========================================================
# OPENAI
# =========================================================

def translate_openai(
    text,
    target_language,
    model,
    api_key,
):

    client = get_openai_client(api_key)

    response = client.responses.create(
        model=model,
        input=[
            {
                "role": "system",
                "content": SYSTEM_PROMPT,
            },
            {
                "role": "user",
                "content": (
                    f"Target language: {target_language}\n\n"
                    f"TEXT:\n{text}"
                ),
            },
        ],
    )

    result = response.output_text

    if not result:
        raise RuntimeError(
            "OpenAI returned an empty response."
        )

    return result.strip()


# =========================================================
# PROVIDER ROUTER
# =========================================================

def translate_once(
    text,
    target_language,
    provider_name,
):

    provider = PROVIDERS.get(
        provider_name
    )

    if not provider:
        raise RuntimeError(
            f"Provider '{provider_name}' is not configured."
        )

    if provider["type"] == "gemini":

        return translate_gemini(
            text,
            target_language,
            provider["model"],
            provider["key"],
        )

    if provider["type"] == "openai":

        return translate_openai(
            text,
            target_language,
            provider["model"],
            provider["key"],
        )

    raise RuntimeError(
        "Unknown provider."
    )


# =========================================================
# RETRY
# =========================================================

def translate_with_retry(
    text,
    target_language,
    provider_name,
    retries=3,
):

    last_error = None

    for attempt in range(retries):

        try:

            return translate_once(
                text,
                target_language,
                provider_name,
            )

        except Exception as error:

            last_error = error

            if attempt < retries - 1:

                wait_time = 2 ** attempt

                time.sleep(
                    wait_time
                )

    raise RuntimeError(
        f"Translation failed after retries: {last_error}"
    )


# =========================================================
# CHUNKING
# =========================================================

def split_text(
    text,
    max_chars=6500,
):

    text = text.replace(
        "\r\n",
        "\n",
    )

    paragraphs = text.split(
        "\n"
    )

    chunks = []
    current = ""

    for paragraph in paragraphs:

        paragraph = paragraph.strip()

        if not paragraph:
            continue

        if (
            len(current)
            + len(paragraph)
            + 2
            <= max_chars
        ):

            if current:
                current += "\n\n"

            current += paragraph

        else:

            if current:
                chunks.append(
                    current
                )

            # Handle extremely long paragraphs
            if len(paragraph) > max_chars:

                for i in range(
                    0,
                    len(paragraph),
                    max_chars,
                ):

                    chunks.append(
                        paragraph[
                            i:i + max_chars
                        ]
                    )

                current = ""

            else:

                current = paragraph

    if current:
        chunks.append(
            current
        )

    return chunks


# =========================================================
# PARALLEL TRANSLATION
# =========================================================

def translate_chunks_parallel(
    chunks,
    target_language,
    provider_name,
    workers=4,
):

    if not chunks:
        return []

    results = [
        None
        for _ in chunks
    ]

    progress = st.progress(
        0,
        text="Preparing translation..."
    )

    status = st.empty()

    completed = 0

    # Controlled parallelism.
    # This improves speed without firing
    # hundreds of requests at once.
    with ThreadPoolExecutor(
        max_workers=workers
    ) as executor:

        future_map = {}

        for index, chunk in enumerate(
            chunks
        ):

            future = executor.submit(
                translate_with_retry,
                chunk,
                target_language,
                provider_name,
            )

            future_map[
                future
            ] = index

        for future in as_completed(
            future_map
        ):

            index = future_map[
                future
            ]

            results[index] = future.result()

            completed += 1

            progress.progress(
                completed / len(chunks),
                text=(
                    f"Translating "
                    f"{completed}/{len(chunks)}"
                ),
            )

            status.write(
                f"Completed {completed} of {len(chunks)} sections"
            )

    return results


# =========================================================
# TEXT TRANSLATION
# =========================================================

def translate_long_text(
    text,
    target_language,
    provider_name,
    workers=4,
):

    chunks = split_text(
        text,
        max_chars=6500,
    )

    if not chunks:
        return ""

    translated = translate_chunks_parallel(
        chunks,
        target_language,
        provider_name,
        workers=workers,
    )

    return "\n\n".join(
        translated
    )


# =========================================================
# TXT
# =========================================================

def extract_text_file(data):

    return data.decode(
        "utf-8",
        errors="replace",
    )


# =========================================================
# DOCX
# =========================================================

def extract_docx(data):

    if Document is None:
        raise RuntimeError(
            "python-docx is not installed."
        )

    document = Document(
        io.BytesIO(data)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            paragraphs.append(
                paragraph.text
            )

    return "\n\n".join(
        paragraphs
    )


# =========================================================
# CSV
# =========================================================

def extract_csv(data):

    text = data.decode(
        "utf-8",
        errors="replace",
    )

    reader = csv.reader(
        io.StringIO(text)
    )

    rows = []

    for row in reader:

        rows.append(
            " | ".join(row)
        )

    return "\n".join(rows)


# =========================================================
# PDF TEXT
# =========================================================

def extract_pdf_text(data):

    if fitz is None:
        raise RuntimeError(
            "PyMuPDF is not installed."
        )

    document = fitz.open(
        stream=data,
        filetype="pdf",
    )

    pages = []

    for page in document:

        pages.append(
            page.get_text(
                "text"
            )
        )

    document.close()

    return "\n\n".join(
        pages
    )


# =========================================================
# PDF TRANSLATION
# =========================================================

def translate_pdf(
    data,
    target_language,
    provider_name,
    workers=4,
):

    if fitz is None:
        raise RuntimeError(
            "PyMuPDF is not installed."
        )

    document = fitz.open(
        stream=data,
        filetype="pdf",
    )

    total_pages = len(
        document
    )

    progress = st.progress(
        0,
        text="Preparing PDF..."
    )

    for page_index in range(
        total_pages
    ):

        page = document[
            page_index
        ]

        text = page.get_text(
            "text"
        ).strip()

        if not text:

            progress.progress(
                (page_index + 1)
                / total_pages,
                text=(
                    f"Page "
                    f"{page_index + 1}/"
                    f"{total_pages}"
                ),
            )

            continue

        chunks = split_text(
            text,
            max_chars=6500,
        )

        translated_chunks = (
            translate_chunks_parallel(
                chunks,
                target_language,
                provider_name,
                workers,
            )
        )

        translated_text = (
            "\n\n".join(
                translated_chunks
            )
        )

        # Detect text blocks.
        blocks = page.get_text(
            "blocks"
        )

        text_rects = []

        for block in blocks:

            if len(block) < 5:
                continue

            block_text = str(
                block[4]
            ).strip()

            if block_text:

                text_rects.append(
                    fitz.Rect(
                        block[:4]
                    )
                )

        if text_rects:

            # Remove only detected text areas.
            for rect in text_rects:

                page.add_redact_annot(
                    rect,
                    fill=(1, 1, 1),
                )

            page.apply_redactions()

            # Reinsert translated text.
            union_rect = text_rects[0]

            for rect in text_rects[1:]:

                union_rect |= rect

            union_rect.x0 += 3
            union_rect.y0 += 3
            union_rect.x1 -= 3
            union_rect.y1 -= 3

            direction = (
                "rtl"
                if is_rtl(
                    target_language
                )
                else "ltr"
            )

            alignment = (
                "right"
                if direction == "rtl"
                else "left"
            )

            safe_text = html.escape(
                translated_text
            ).replace(
                "\n",
                "<br>",
            )

            html_content = f"""
            <div
                dir="{direction}"
                style="
                    font-family: sans-serif;
                    font-size: 10pt;
                    line-height: 1.35;
                    text-align: {alignment};
                "
            >
                {safe_text}
            </div>
            """

            try:

                page.insert_htmlbox(
                    union_rect,
                    html_content,
                    css="",
                )

            except Exception:

                page.insert_textbox(
                    union_rect,
                    translated_text,
                    fontsize=9,
                    fontname="helv",
                    color=(0, 0, 0),
                )

        progress.progress(
            (page_index + 1)
            / total_pages,
            text=(
                f"Page "
                f"{page_index + 1}/"
                f"{total_pages}"
            ),
        )

    output = document.tobytes(
        garbage=4,
        deflate=True,
    )

    document.close()

    return output


# =========================================================
# SESSION DEFAULTS
# =========================================================

if "provider" not in st.session_state:

    if "Gemini" in PROVIDERS:
        st.session_state.provider = "Gemini"

    elif PROVIDERS:
        st.session_state.provider = next(
            iter(PROVIDERS)
        )

    else:
        st.session_state.provider = None


if "workers" not in st.session_state:
    st.session_state.workers = 4


# =========================================================
# SIDEBAR SETTINGS
# =========================================================

with st.sidebar:

    st.markdown("## ⚙ Settings")

    st.markdown(
        "### Translation Engine"
    )

    if PROVIDERS:

        provider_names = list(
            PROVIDERS.keys()
        )

        current = st.session_state.provider

        if current not in provider_names:
            current = provider_names[0]

        selected_provider = st.selectbox(
            "Provider",
            provider_names,
            index=provider_names.index(
                current
            ),
        )

        st.session_state.provider = (
            selected_provider
        )

        provider_info = PROVIDERS[
            selected_provider
        ]

        st.caption(
            f"Model: {provider_info['model']}"
        )

        st.markdown(
            '<span class="status">'
            '<span class="dot"></span>'
            ' Connected'
            '</span>',
            unsafe_allow_html=True,
        )

    else:

        st.error(
            "No translation provider is configured."
        )

    st.markdown("---")

    st.markdown(
        "### Performance"
    )

    workers = st.slider(
        "Parallel requests",
        min_value=1,
        max_value=8,
        value=st.session_state.workers,
        help=(
            "Higher values can be faster, "
            "but may hit API rate limits."
        ),
    )

    st.session_state.workers = workers

    st.markdown("---")

    st.markdown(
        """
        <div class="small-muted">
        API keys are loaded from Streamlit Secrets.
        They are not displayed in the interface.
        </div>
        """,
        unsafe_allow_html=True,
    )


# =========================================================
# MAIN CONTROLS
# =========================================================

col1, col2 = st.columns(
    [1, 1]
)

with col1:

    target_language = st.selectbox(
        "Translate to",
        LANGUAGES,
        index=0,
    )

with col2:

    if st.session_state.provider:

        st.selectbox(
            "Engine",
            [
                st.session_state.provider
            ],
            disabled=True,
        )


st.markdown(
    '<div class="glass">',
    unsafe_allow_html=True,
)


input_mode = st.radio(
    "Input type",
    ["Text", "File"],
    horizontal=True,
    label_visibility="collapsed",
)


# =========================================================
# TEXT MODE
# =========================================================

if input_mode == "Text":

    text = st.text_area(
        "Text",
        height=270,
        placeholder=(
            "Write or paste your text here..."
        ),
        label_visibility="collapsed",
    )

    if st.button(
        "✦  Translate",
        type="primary",
        use_container_width=True,
    ):

        if not PROVIDERS:

            st.error(
                "No API provider is configured."
            )

        elif not text.strip():

            st.warning(
                "Enter some text first."
            )

        else:

            try:

                result = translate_long_text(
                    text,
                    target_language,
                    st.session_state.provider,
                    st.session_state.workers,
                )

                st.success(
                    "Translation completed."
                )

                st.text_area(
                    "Result",
                    result,
                    height=400,
                )

                st.download_button(
                    "Download TXT",
                    result.encode(
                        "utf-8"
                    ),
                    file_name=(
                        "translation.txt"
                    ),
                    mime="text/plain",
                    use_container_width=True,
                )

            except Exception as error:

                st.error(
                    f"Translation failed: {error}"
                )


# =========================================================
# FILE MODE
# =========================================================

else:

    uploaded = st.file_uploader(
        "Upload document",
        type=[
            "pdf",
            "docx",
            "txt",
            "md",
            "csv",
        ],
        label_visibility="collapsed",
    )

    if uploaded:

        file_name = uploaded.name

        extension = Path(
            file_name
        ).suffix.lower()

        st.markdown(
            f"**{file_name}**"
        )

        if st.button(
            "✦  Translate document",
            type="primary",
            use_container_width=True,
        ):

            if not PROVIDERS:

                st.error(
                    "No API provider is configured."
                )

            else:

                try:

                    data = uploaded.read()

                    provider = (
                        st.session_state.provider
                    )

                    workers = (
                        st.session_state.workers
                    )

                    # -----------------------------
                    # PDF
                    # -----------------------------

                    if extension == ".pdf":

                        with st.spinner(
                            "Processing document..."
                        ):

                            result = translate_pdf(
                                data,
                                target_language,
                                provider,
                                workers,
                            )

                        st.success(
                            "PDF translation completed."
                        )

                        st.download_button(
                            "Download translated PDF",
                            result,
                            file_name=(
                                Path(
                                    file_name
                                ).stem
                                + "_translated.pdf"
                            ),
                            mime="application/pdf",
                            use_container_width=True,
                        )

                    # -----------------------------
                    # DOCX
                    # -----------------------------

                    elif extension == ".docx":

                        text = extract_docx(
                            data
                        )

                        result = (
                            translate_long_text(
                                text,
                                target_language,
                                provider,
                                workers,
                            )
                        )

                        st.success(
                            "Translation completed."
                        )

                        st.text_area(
                            "Translation",
                            result,
                            height=450,
                        )

                        st.download_button(
                            "Download TXT",
                            result.encode(
                                "utf-8"
                            ),
                            file_name=(
                                Path(
                                    file_name
                                ).stem
                                + "_translated.txt"
                            ),
                            mime="text/plain",
                            use_container_width=True,
                        )

                    # -----------------------------
                    # TXT / MD
                    # -----------------------------

                    elif extension in {
                        ".txt",
                        ".md",
                    }:

                        text = (
                            extract_text_file(
                                data
                            )
                        )

                        result = (
                            translate_long_text(
                                text,
                                target_language,
                                provider,
                                workers,
                            )
                        )

                        st.success(
                            "Translation completed."
                        )

                        st.text_area(
                            "Translation",
                            result,
                            height=450,
                        )

                        st.download_button(
                            "Download TXT",
                            result.encode(
                                "utf-8"
                            ),
                            file_name=(
                                Path(
                                    file_name
                                ).stem
                                + "_translated.txt"
                            ),
                            mime="text/plain",
                            use_container_width=True,
                        )

                    # -----------------------------
                    # CSV
                    # -----------------------------

                    elif extension == ".csv":

                        text = extract_csv(
                            data
                        )

                        result = (
                            translate_long_text(
                                text,
                                target_language,
                                provider,
                                workers,
                            )
                        )

                        st.success(
                            "Translation completed."
                        )

                        st.text_area(
                            "Translation",
                            result,
                            height=450,
                        )

                        st.download_button(
                            "Download TXT",
                            result.encode(
                                "utf-8"
                            ),
                            file_name=(
                                Path(
                                    file_name
                                ).stem
                                + "_translated.txt"
                            ),
                            mime="text/plain",
                            use_container_width=True,
                        )

                except Exception as error:

                    st.error(
                        f"Translation failed: {error}"
                    )


st.markdown(
    "</div>",
    unsafe_allow_html=True,
)


# =========================================================
# FOOTER
# =========================================================

st.markdown(
    """
    <div style="
        text-align:center;
        color:#688383;
        font-size:12px;
        padding-top:25px;
    ">
        Translate Preserv · Smart document translation
    </div>
    """,
    unsafe_allow_html=True,
              )
